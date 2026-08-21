import gzip
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from benchmark_core import (
    BenchmarkEvent,
    BenchmarkPlan,
    PlanType,
    RequestError,
    RequestSample,
    TaskPhase,
    TimingMetrics,
    TokenSource,
    TokenUsage,
    aggregate_time_windows,
    build_summary,
    write_report_bundle,
)
from benchmark_core.models import ErrorCategory


class ReportingTests(unittest.TestCase):
    def setUp(self):
        self.plan = BenchmarkPlan("报告测试", PlanType.FIXED_CONCURRENCY, concurrency=2, total_requests=2)
        self.samples = (
            RequestSample(
                "request-1",
                0.0,
                TimingMetrics(1.2, 0.2, 1.0),
                TokenUsage(5, 10, 15, TokenSource.SERVER),
                content="ok",
                finish_reason="stop",
                transport_success=True,
                protocol_valid=True,
                assertion_passed=True,
                output_tps=10.0,
            ),
            RequestSample(
                "request-2",
                0.2,
                TimingMetrics(2.0, 0.4, 1.6),
                TokenUsage(source=TokenSource.UNAVAILABLE),
                error=RequestError(ErrorCategory.SERVER, "Bearer abcdefghijk sk-abcdefghijk"),
            ),
        )
        self.summary = build_summary("task-report", self.plan, self.samples, 2.5)
        self.events = (BenchmarkEvent(1, 0.0, TaskPhase.LOAD, "phase_started"),)
        self.windows = aggregate_time_windows(self.samples, total_elapsed=2.5, window_seconds=1.0)

    def test_bundle_formats_are_readable_hashed_and_secret_free(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            artifacts = write_report_bundle(
                root,
                self.summary,
                self.samples,
                self.events,
                self.windows,
                ["json", "jsonl.gz", "events.jsonl.gz", "csv", "xlsx", "html"],
            )
            self.assertEqual(6, len(artifacts))
            for artifact in artifacts:
                self.assertEqual(Path(artifact.relative_path).name, artifact.relative_path)
                path = root / artifact.relative_path
                raw = path.read_bytes()
                self.assertEqual(len(raw), artifact.size_bytes)
                self.assertEqual(hashlib.sha256(raw).hexdigest(), artifact.sha256)
                self.assertNotIn(b"abcdefghijk", raw)

            payload = json.loads((root / "task-report.json").read_text(encoding="utf-8"))
            self.assertEqual("1.0", payload["schema_version"])
            self.assertEqual("fixed_concurrency", payload["plan"]["plan_type"])
            self.assertIn("windows", payload)

            with gzip.open(root / "task-report.samples.jsonl.gz", "rt", encoding="utf-8") as handle:
                sample_rows = [json.loads(line) for line in handle]
            self.assertEqual(2, len(sample_rows))
            self.assertEqual("request-1", sample_rows[0]["request_id"])

            with gzip.open(root / "task-report.events.jsonl.gz", "rt", encoding="utf-8") as handle:
                event_rows = [json.loads(line) for line in handle]
            self.assertEqual("phase_started", event_rows[0]["event_type"])

            self.assertTrue((root / "task-report.csv").read_bytes().startswith(b"\xef\xbb\xbf"))
            workbook = load_workbook(root / "task-report.xlsx", read_only=True)
            self.assertEqual(["摘要", "请求样本"], workbook.sheetnames)
            workbook.close()

            document = (root / "task-report.html").read_text(encoding="utf-8")
            self.assertIn("<!doctype html>", document.lower())
            self.assertIn("const D=", document)
            self.assertNotIn("http://", document)
            self.assertNotIn("https://", document)
            self.assertNotIn("src=", document.lower())
            self.assertNotIn("href=", document.lower())


if __name__ == "__main__":
    unittest.main()
